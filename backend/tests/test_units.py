"""Pure unit tests (no HTTP)."""

import io
from datetime import date, timedelta
from decimal import Decimal as D
from types import SimpleNamespace

import pytest

from app.ai.assistant import validate_citations
from app.ai.intents import classify
from app.ai.retrieval import tokenize
from app.models.enums import JobStatus, OrderStatus
from app.permissions.policy import P, Role, has_permission
from app.services.extraction import extract
from app.services.orders import calculate_totals, can_transition
from app.services.production import is_overdue


def test_order_totals_rounding_and_tax():
    lines, sub, tax, total = calculate_totals([(D("500"), D("4.20")), (D("3"), D("0.335"))], D("20"))
    assert lines == [D("2100.00"), D("1.01")]
    assert sub == D("2101.01") and tax == D("420.20") and total == D("2521.21")


def test_order_totals_zero_tax():
    _, sub, tax, total = calculate_totals([(D("1.5"), D("10"))], D("0"))
    assert (sub, tax, total) == (D("15.00"), D("0.00"), D("15.00"))


@pytest.mark.parametrize(
    "current,target,ok",
    [
        (OrderStatus.DRAFT, OrderStatus.CONFIRMED, True),
        (OrderStatus.DRAFT, OrderStatus.FULFILLED, False),
        (OrderStatus.CONFIRMED, OrderStatus.IN_PRODUCTION, True),
        (OrderStatus.CONFIRMED, OrderStatus.DRAFT, False),
        (OrderStatus.IN_PRODUCTION, OrderStatus.PARTIALLY_FULFILLED, True),
        (OrderStatus.PARTIALLY_FULFILLED, OrderStatus.FULFILLED, True),
        (OrderStatus.PARTIALLY_FULFILLED, OrderStatus.CANCELLED, False),
        (OrderStatus.FULFILLED, OrderStatus.CANCELLED, False),
        (OrderStatus.CANCELLED, OrderStatus.CONFIRMED, False),
    ],
)
def test_order_state_machine(current, target, ok):
    assert can_transition(current, target) is ok


@pytest.mark.parametrize(
    "role,perm,ok",
    [
        (Role.SALES, P.ORDERS_WRITE, True),
        (Role.SALES, P.INVENTORY_ADJUST, False),
        (Role.SALES, P.USERS_MANAGE, False),
        (Role.INVENTORY_MANAGER, P.INVENTORY_ADJUST, True),
        (Role.INVENTORY_MANAGER, P.ORDERS_WRITE, False),
        (Role.PRODUCTION_MANAGER, P.PRODUCTION_WRITE, True),
        (Role.PRODUCTION_MANAGER, P.INVENTORY_ADJUST, False),
        (Role.QUALITY_INSPECTOR, P.QUALITY_WRITE, True),
        (Role.QUALITY_INSPECTOR, P.ORDERS_WRITE, False),
        (Role.DISPATCH, P.DELIVERIES_WRITE, True),
        (Role.DISPATCH, P.PRODUCTION_WRITE, False),
        (Role.MANAGEMENT, P.DASHBOARD_READ, True),
        (Role.MANAGEMENT, P.ORDERS_WRITE, False),
        (Role.ADMIN, P.AUDIT_READ, True),
        ("NOT_A_ROLE", P.DASHBOARD_READ, False),
    ],
)
def test_permission_matrix(role, perm, ok):
    assert has_permission(role, perm) is ok


def test_management_is_read_only():
    from app.permissions.policy import ROLE_PERMISSIONS

    assert all(p.value.endswith((":read", ":use")) for p in ROLE_PERMISSIONS[Role.MANAGEMENT])


def test_overdue_detection():
    today = date(2026, 10, 19)
    job = SimpleNamespace(status=JobStatus.IN_PROGRESS, target_date=today - timedelta(days=1))
    assert is_overdue(job, today)
    job.status = JobStatus.COMPLETED
    assert not is_overdue(job, today)
    assert not is_overdue(SimpleNamespace(status=JobStatus.PLANNED, target_date=None), today)


def test_intent_routing():
    assert classify("Summarize order SO-1001 and its current status.").params == {"order_number": "SO-1001"}
    assert classify("Which materials are below their reorder level?").name == "low_stock"
    assert classify("Which production jobs are overdue?").name == "overdue_jobs"
    assert classify("Which orders are waiting for quality inspection?").name == "awaiting_inspection"
    assert classify("What does the steel specification say about grade?").name == "documents"


def test_tokenizer_splits_hyphens_keeps_decimals():
    assert tokenize("Quality-control tolerance 0.5 mm, the S275JR") == ["quality", "control", "tolerance", "0.5", "mm", "s275jr"]


def test_citation_validation_strips_fabricated_sources():
    text, used = validate_citations("Grade is S275JR [1]. Also see [7] and [2].", 2)
    assert used == [1, 2]
    assert "[7]" not in text


def test_markdown_extraction_keeps_sections():
    ex = extract(b"# Title\n\n## Sampling\nInspect 20 units.\n\n## Records\nKeep 7 years.", ".md")
    assert [(c.section, c.text) for c in ex.chunks] == [("Sampling", "Inspect 20 units."), ("Records", "Keep 7 years.")]


def test_docx_extraction_uses_headings():
    import docx

    d = docx.Document()
    d.add_heading("Welding", level=1)
    d.add_paragraph("Preheat plate above 25 mm.")
    buf = io.BytesIO()
    d.save(buf)
    ex = extract(buf.getvalue(), ".docx")
    assert ex.chunks[0].section == "Welding" and "Preheat" in ex.chunks[0].text


def test_pdf_extraction_page_numbers():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(200, 200)
    buf = io.BytesIO()
    w.write(buf)
    ex = extract(buf.getvalue(), ".pdf")
    assert ex.page_count == 1 and ex.chunks == [] and "OCR" in ex.note
