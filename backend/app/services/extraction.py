"""Text extraction and chunking for business documents.

PDF: text layer only, one or more chunks per page (scanned PDFs without a text layer are
flagged; OCR is not included). DOCX: paragraphs, with Heading styles as section names.
Markdown: '#' headings as sections. TXT: paragraphs.
"""

import io
import re
from dataclasses import dataclass

MAX_CHARS = 900


@dataclass
class Chunk:
    text: str
    page_number: int | None = None
    section: str | None = None


@dataclass
class Extraction:
    chunks: list[Chunk]
    page_count: int | None = None
    note: str | None = None


def _split(text: str, page: int | None, section: str | None) -> list[Chunk]:
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        p = re.sub(r"[ \t]+", " ", p)
        if buf and len(buf) + len(p) + 1 > MAX_CHARS:
            chunks.append(Chunk(buf, page, section))
            buf = ""
        while len(p) > MAX_CHARS:  # very long paragraph: hard split on sentence/space boundary
            cut = p.rfind(". ", 0, MAX_CHARS)
            cut = cut + 1 if cut > MAX_CHARS // 2 else (p.rfind(" ", 0, MAX_CHARS) or MAX_CHARS)
            chunks.append(Chunk(p[:cut].strip(), page, section))
            p = p[cut:].strip()
        buf = f"{buf}\n{p}" if buf else p
    if buf:
        chunks.append(Chunk(buf, page, section))
    return chunks


def _sections(lines: list[tuple[str | None, str]]) -> list[Chunk]:
    """lines: (heading_or_None, text). Group text under the most recent heading."""
    out, section, buf = [], None, []
    for heading, text in lines:
        if heading is not None:
            if buf:
                out += _split("\n\n".join(buf), None, section)
            section, buf = heading[:200], []
        elif text.strip():
            buf.append(text)
    if buf:
        out += _split("\n\n".join(buf), None, section)
    return out


def extract_pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    chunks: list[Chunk] = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            chunks += _split(text, i, None)
    note = None
    if not chunks:
        note = (
            "No extractable text layer found (possibly a scanned PDF). OCR is not enabled, so this document cannot be searched."
        )
    return Extraction(chunks, len(reader.pages), note)


def extract_docx(data: bytes) -> Extraction:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines = []
    for p in d.paragraphs:
        style = (p.style.name or "") if p.style is not None else ""
        if style.lower().startswith("heading") or style.lower() == "title":
            lines.append((p.text.strip() or None, ""))
        else:
            lines.append((None, p.text))
    for t in d.tables:
        for row in t.rows:
            lines.append((None, " | ".join(c.text.strip() for c in row.cells)))
    return Extraction(_sections(lines))


def extract_markdown(text: str) -> Extraction:
    lines = []
    for line in text.splitlines():
        m = re.match(r"^\s{0,3}#{1,6}\s+(.*)$", line)
        lines.append((m.group(1).strip(), "") if m else (None, line))
    # re-join consecutive non-heading lines into paragraphs
    merged, buf = [], []
    for h, t in lines:
        if h is not None:
            if buf:
                merged.append((None, "\n".join(buf)))
                buf = []
            merged.append((h, ""))
        elif t.strip() == "":
            if buf:
                merged.append((None, "\n".join(buf)))
                buf = []
        else:
            buf.append(t)
    if buf:
        merged.append((None, "\n".join(buf)))
    return Extraction(_sections(merged))


def extract(data: bytes, ext: str) -> Extraction:
    if ext == ".pdf":
        return extract_pdf(data)
    if ext == ".docx":
        return extract_docx(data)
    text = data.decode("utf-8", errors="replace")
    if ext in (".md", ".markdown"):
        return extract_markdown(text)
    return Extraction(_split(text, None, None))
